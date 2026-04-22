from __future__ import annotations

import copy
import io
import os
import sys
from dataclasses import dataclass, field
from typing import List, Tuple

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

if sys.platform == "win32":
    sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8")
    sys.stdin = io.TextIOWrapper(sys.stdin.buffer, encoding="utf-8")

SECTION_MAP = {
    "p": "前奏",
    "a": "A melo",
    "b": "B melo",
    "c": "C melo",
    "r": "副歌",
    "i": "间奏",
    "o": "尾奏",
}
PURE_ACTION_LYRIC = ("（纯动作/无歌词）", "（纯动作/无歌词）")


@dataclass
class Block:
    block_type: str
    beats: str
    lyrics: List[Tuple[str, str]]


@dataclass
class ArrangementState:
    idx: int = 0
    current_lyrics: List[Tuple[str, str]] = field(default_factory=list)
    blocks: List[Block] = field(default_factory=list)


class WotaArrangementTool:
    def __init__(self) -> None:
        self.song_name = ""
        self.bpm = ""
        self.lyric_pairs: List[Tuple[str, str]] = []
        self.state = ArrangementState()
        self.history: List[ArrangementState] = []

    def run(self) -> None:
        self._print_banner()
        self._load_inputs()
        self._run_intro_phase()
        self._run_lyric_pipeline()
        self._run_tail_pack_phase()
        self._run_outro_phase()
        self._render_excel()

    @staticmethod
    def _print_banner() -> None:
        print("\n" + "=" * 60)
        print("   🌟  Wota-Arrangement-Tool v1.0 🌟")
        print("=" * 60)

    def _load_inputs(self) -> None:
        lines: List[str] = []
        while True:
            filepath = input("\n 📂 1. 请输入歌词txt路径: ").strip().strip('"').strip("'")
            if os.path.exists(filepath):
                with open(filepath, "r", encoding="utf-8") as file:
                    lines = [line.strip() for line in file if line.strip()]
                if len(lines) >= 2:
                    break
            print(" ❌ 路径无效或内容为空，请重试。")

        self.song_name = input(" 🎵 2. 请输入歌曲名称: ").strip() or "Untitled"
        self.bpm = input(" ⚡ 3. 请输入歌曲 BPM: ").strip() or "120"
        self.lyric_pairs = [(lines[i], lines[i + 1]) for i in range(0, len(lines) - 1, 2)]

    def _save_history(self) -> None:
        self.history.append(copy.deepcopy(self.state))

    def _undo(self) -> bool:
        if self.history:
            self.state = self.history.pop()
            print("\n   ↩️  撤销成功，已回退到上一步状态")
            return True
        print("\n   ⚠️  已经是最初状态，无法撤销")
        return False

    @staticmethod
    def _resolve_section(block_code: str) -> str:
        return SECTION_MAP.get(block_code, block_code)

    def _append_pure_action_block(self, block_type: str, beats: str) -> None:
        self.state.blocks.append(Block(block_type=block_type, beats=beats, lyrics=[PURE_ACTION_LYRIC]))

    def _run_intro_phase(self) -> None:
        print("\n" + "─" * 30 + "\n 🎹 [前奏阶段] 输入 'p 拍数'，回车进入歌词")
        while True:
            intro_count = sum(1 for block in self.state.blocks if block.block_type == "前奏")
            cmd = input(f" [已记录 {intro_count} 段前奏] >> ").strip().lower()
            if not cmd:
                return
            if cmd == "u":
                self._undo()
                continue

            parts = cmd.split()
            if len(parts) >= 2 and parts[0] == "p":
                self._save_history()
                self._append_pure_action_block("前奏", parts[1])
            else:
                print("   ⚠️ 格式错误，例如 'p 8'")

    def _run_lyric_pipeline(self) -> None:
        print("\n" + "═" * 60 + "\n 【 歌词处理中 】 回车收纳 / 代码打包 / u 撤销\n" + "═" * 60)

        while self.state.idx < len(self.lyric_pairs):
            jp, cn = self.lyric_pairs[self.state.idx]
            cache_info = f"📦 暂存: {len(self.state.current_lyrics)}句" if self.state.current_lyrics else ""
            print(f"\n 🟢 [进度: {self.state.idx + 1}/{len(self.lyric_pairs)}] {cache_info}")
            print(f" --------------------------------------------------\n  日: {jp}\n  中: {cn}\n --------------------------------------------------")
            cmd = input(" >> 指令: ").strip().lower()

            if cmd == "u":
                self._undo()
                continue

            if cmd == "":
                self._save_history()
                self.state.current_lyrics.append((jp, cn))
                self.state.idx += 1
                continue

            parts = cmd.split()
            if len(parts) < 2:
                print(" ⚠️ 格式不对（例：r 16）")
                continue

            self._save_history()
            block_code, beats = parts[0], parts[1]
            if block_code == "i" and not self.state.current_lyrics:
                self._append_pure_action_block("间奏", beats)
                continue

            self.state.current_lyrics.append((jp, cn))
            self.state.blocks.append(
                Block(
                    block_type=self._resolve_section(block_code),
                    beats=beats,
                    lyrics=self.state.current_lyrics.copy(),
                )
            )
            self.state.current_lyrics = []
            self.state.idx += 1

    def _run_tail_pack_phase(self) -> None:
        while self.state.current_lyrics:
            print(f"\n 📢 剩余 {len(self.state.current_lyrics)} 句歌词，请输入 [类型 拍数] 打包，或 'u' 撤销:")
            cmd = input(" >> ").strip().lower()
            if cmd == "u":
                self._undo()
                if self.state.idx < len(self.lyric_pairs):
                    break
                continue

            parts = cmd.split()
            if len(parts) < 2:
                print(" ⚠️ 格式不对（例：r 16）")
                continue

            self._save_history()
            self.state.blocks.append(
                Block(
                    block_type=self._resolve_section(parts[0]),
                    beats=parts[1],
                    lyrics=self.state.current_lyrics.copy(),
                )
            )
            self.state.current_lyrics = []

    def _run_outro_phase(self) -> None:
        print("\n" + "─" * 30 + "\n 🎸 [尾奏阶段] 输入 'o 拍数'，直接回车保存")
        while True:
            outro_count = sum(1 for block in self.state.blocks if block.block_type == "尾奏")
            cmd = input(f" [已记录 {outro_count} 段尾奏] >> ").strip().lower()
            if not cmd:
                return
            if cmd == "u":
                self._undo()
                continue

            parts = cmd.split()
            if len(parts) >= 2 and parts[0] == "o":
                self._save_history()
                self._append_pure_action_block("尾奏", parts[1])
            else:
                print("   ⚠️ 格式错误，例如 'o 4'")

    @staticmethod
    def _sanitize_output_name(song_name: str) -> str:
        output_name = f"{song_name}_编排表.xlsx"
        for char in ['\\', "/", ":", "*", "?", '"', "<", ">", "|"]:
            output_name = output_name.replace(char, "_")
        return output_name

    def _render_excel(self) -> None:
        output_name = self._sanitize_output_name(self.song_name)
        while True:
            try:
                if os.path.exists(output_name):
                    os.remove(output_name)
                workbook = Workbook()
                worksheet = workbook.active
                worksheet.title = "打艺编排脚本"

                fill_odd = PatternFill(start_color="F9FAFB", end_color="F9FAFB", fill_type="solid")
                fill_even = PatternFill(start_color="EFF3F6", end_color="EFF3F6", fill_type="solid")
                header_fill = PatternFill(start_color="1F2D3D", end_color="1F2D3D", fill_type="solid")
                border = Border(
                    left=Side(style="thin", color="DCDFE6"),
                    right=Side(style="thin", color="DCDFE6"),
                    top=Side(style="thin", color="DCDFE6"),
                    bottom=Side(style="thin", color="DCDFE6"),
                )

                worksheet.merge_cells("A1:F1")
                worksheet["A1"] = f"{self.song_name} (BPM: {self.bpm})"
                worksheet["A1"].alignment = Alignment(horizontal="center", vertical="center")
                worksheet["A1"].font = Font(size=18, bold=True)

                headers = ["段落", "拍数", "日文歌词", "中文歌词", "技 / 动作编排", "备注"]
                for col, header in enumerate(headers, 1):
                    cell = worksheet.cell(row=2, column=col, value=header)
                    cell.fill = header_fill
                    cell.font = Font(color="FFFFFF", bold=True)
                    cell.alignment = Alignment(horizontal="center", vertical="center")

                curr_row = 3
                block_row_info: List[Tuple[int, int]] = []
                last_color = fill_even
                for idx, block in enumerate(self.state.blocks):
                    same_as_prev = idx > 0 and block.block_type == self.state.blocks[idx - 1].block_type
                    color = last_color if same_as_prev else (fill_odd if last_color == fill_even else fill_even)
                    last_color = color
                    start_row = curr_row
                    for jp, cn in block.lyrics:
                        worksheet.cell(row=curr_row, column=3, value=jp)
                        worksheet.cell(row=curr_row, column=4, value=cn)
                        for col in range(1, 7):
                            cell = worksheet.cell(row=curr_row, column=col)
                            cell.fill = color
                            cell.border = border
                            cell.alignment = Alignment(vertical="center", wrap_text=True, indent=1)
                        worksheet.row_dimensions[curr_row].height = 30
                        curr_row += 1
                    block_row_info.append((start_row, curr_row - 1))

                for idx, block in enumerate(self.state.blocks):
                    start_row, end_row = block_row_info[idx]
                    for col in (2, 5, 6):
                        if end_row > start_row:
                            worksheet.merge_cells(start_row=start_row, start_column=col, end_row=end_row, end_column=col)
                        if col == 2:
                            worksheet.cell(row=start_row, column=col, value=block.beats).alignment = Alignment(
                                horizontal="center", vertical="center"
                            )

                i = 0
                while i < len(self.state.blocks):
                    j = i
                    while j + 1 < len(self.state.blocks) and self.state.blocks[j + 1].block_type == self.state.blocks[i].block_type:
                        j += 1
                    start_merge = block_row_info[i][0]
                    end_merge = block_row_info[j][1]
                    if end_merge > start_merge:
                        worksheet.merge_cells(start_row=start_merge, start_column=1, end_row=end_merge, end_column=1)
                    section_cell = worksheet.cell(row=start_merge, column=1, value=self.state.blocks[i].block_type)
                    section_cell.font = Font(bold=True, color="409EFF")
                    section_cell.alignment = Alignment(horizontal="center", vertical="center")
                    i = j + 1

                worksheet.column_dimensions["C"].width = 40
                worksheet.column_dimensions["D"].width = 40
                workbook.save(output_name)
                print(f"\n ✨ 成功！文件已保存至: {output_name}")
                return
            except PermissionError:
                input(f"\n ⚠️  文件 [{output_name}] 被占用，请关闭它后按回车重试...")


def main() -> None:
    tool = WotaArrangementTool()
    tool.run()


if __name__ == "__main__":
    try:
        main()
    except KeyboardInterrupt:
        print("\n 👋 已退出")
