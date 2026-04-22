from __future__ import annotations

from pathlib import Path
from typing import List, Tuple

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

from .models import ArrangementState, Block, HEADERS, PURE_ACTION_LYRIC, TITLE_PATTERN


class XlsxIO:
    song_name: str
    bpm: str
    state: ArrangementState
    default_output_path: "Path | None"

    @staticmethod
    def _cell_str(value: object) -> str:
        if value is None:
            return ""
        return str(value).strip()

    @classmethod
    def _parse_title(cls, raw_title: str) -> tuple[str, str]:
        match = TITLE_PATTERN.match(raw_title.strip())
        if not match:
            return raw_title.strip() or "Untitled", "120"
        return match.group("name").strip() or "Untitled", match.group("bpm").strip() or "120"

    @classmethod
    def _read_state_from_xlsx(cls, path: Path) -> tuple[str, str, ArrangementState]:
        workbook = load_workbook(path)
        worksheet = workbook.active

        actual_headers = [cls._cell_str(worksheet.cell(row=2, column=col).value) for col in range(1, 7)]
        if actual_headers != HEADERS:
            raise ValueError("xlsx 表头不匹配，当前仅支持本工具生成的标准模板。")

        raw_title = cls._cell_str(worksheet["A1"].value)
        song_name, bpm = cls._parse_title(raw_title)

        blocks: List[Block] = []
        current_section = ""
        current_block: Block | None = None

        for row in range(3, worksheet.max_row + 1):
            values = [cls._cell_str(worksheet.cell(row=row, column=col).value) for col in range(1, 7)]
            if not any(values):
                continue

            section, beats, jp, cn, arrangement, remarks = values
            if section:
                current_section = section

            if beats:
                if not current_section:
                    raise ValueError(f"第 {row} 行拍数存在，但段落为空，无法解析。")
                current_block = Block(
                    block_type=current_section,
                    beats=beats,
                    lyrics=[],
                    arrangement=arrangement,
                    remarks=remarks,
                )
                blocks.append(current_block)
            elif current_block is None:
                raise ValueError(f"第 {row} 行无法定位所属段落，xlsx 合并结构可能已损坏。")

            if current_block is not None and (jp or cn):
                current_block.lyrics.append((jp, cn))

        if not blocks:
            raise ValueError("xlsx 中没有可解析的段落数据。")

        for block in blocks:
            if not block.lyrics:
                block.lyrics = [PURE_ACTION_LYRIC]

        return song_name, bpm, ArrangementState(idx=0, current_lyrics=[], blocks=blocks)

    def _load_from_xlsx_prompt(self) -> None:
        while True:
            file_path = input("\n 📂 请输入已生成的编排表 xlsx 路径: ").strip().strip('"').strip("'")
            path = Path(file_path)
            if not path.exists():
                print(" ❌ 文件不存在，请重试。")
                continue
            try:
                self.song_name, self.bpm, self.state = self._read_state_from_xlsx(path)
            except Exception as exc:
                print(f" ❌ 解析失败: {exc}")
                continue
            self.default_output_path = path
            self.history = []
            self.redo_history = []
            print(f" ✅ 已载入：{path.name}（段落 {len(self.state.blocks)} 个）")
            return

    def _build_workbook(self, base_path: "Path | None" = None) -> Workbook:
        if base_path is not None and base_path.exists():
            workbook = load_workbook(base_path)
            worksheet = workbook.active
            # Remove all merged ranges that touch columns A-F
            for rng in list(worksheet.merged_cells.ranges):
                if rng.min_col <= 6:
                    worksheet.unmerge_cells(str(rng))
            # Clear values and styles in columns A-F
            for row in worksheet.iter_rows(min_col=1, max_col=6):
                for cell in row:
                    cell.value = None
                    cell.fill = PatternFill()
                    cell.font = Font()
                    cell.border = Border()
                    cell.alignment = Alignment()
        else:
            workbook = Workbook()
            worksheet = workbook.active
            worksheet.title = "打艺编排脚本"

        fill_odd = PatternFill(start_color="F9FAFB", end_color="F9FAFB", fill_type="solid")
        fill_even = PatternFill(start_color="EFF3F6", end_color="EFF3F6", fill_type="solid")
        header_fill = PatternFill(start_color="1F2D3D", end_color="1F2D3D", fill_type="solid")
        border = Border(
            left=Side(style="thin", color="DCDFE6"),
            right=Side(style="thin", color="DCDFE6"),
            top=Side(style="thin", color="DCDFE6"),
            bottom=Side(style="thin", color="DCDFE6"),
        )

        worksheet.merge_cells("A1:F1")
        worksheet["A1"] = f"{self.song_name} (BPM: {self.bpm})"
        worksheet["A1"].alignment = Alignment(horizontal="center", vertical="center")
        worksheet["A1"].font = Font(size=18, bold=True)

        for col, header in enumerate(HEADERS, 1):
            cell = worksheet.cell(row=2, column=col, value=header)
            cell.fill = header_fill
            cell.font = Font(color="FFFFFF", bold=True)
            cell.alignment = Alignment(horizontal="center", vertical="center")

        curr_row = 3
        block_row_info: List[Tuple[int, int]] = []
        last_color = fill_even
        for idx, block in enumerate(self.state.blocks):
            same_as_prev = idx > 0 and block.block_type == self.state.blocks[idx - 1].block_type
            color = last_color if same_as_prev else (fill_odd if last_color == fill_even else fill_even)
            last_color = color
            start_row = curr_row
            for jp, cn in block.lyrics:
                worksheet.cell(row=curr_row, column=3, value=jp)
                worksheet.cell(row=curr_row, column=4, value=cn)
                for col in range(1, 7):
                    cell = worksheet.cell(row=curr_row, column=col)
                    cell.fill = color
                    cell.border = border
                    cell.alignment = Alignment(vertical="center", wrap_text=True, indent=1)
                worksheet.row_dimensions[curr_row].height = 30
                curr_row += 1
            block_row_info.append((start_row, curr_row - 1))

        for idx, block in enumerate(self.state.blocks):
            start_row, end_row = block_row_info[idx]
            for col in (2, 5, 6):
                if end_row > start_row:
                    worksheet.merge_cells(start_row=start_row, start_column=col, end_row=end_row, end_column=col)
            worksheet.cell(row=start_row, column=2, value=block.beats).alignment = Alignment(horizontal="center", vertical="center")
            worksheet.cell(row=start_row, column=5, value=block.arrangement).alignment = Alignment(horizontal="center", vertical="center")
            worksheet.cell(row=start_row, column=6, value=block.remarks).alignment = Alignment(horizontal="center", vertical="center")

        i = 0
        while i < len(self.state.blocks):
            j = i
            while j + 1 < len(self.state.blocks) and self.state.blocks[j + 1].block_type == self.state.blocks[i].block_type:
                j += 1
            start_merge = block_row_info[i][0]
            end_merge = block_row_info[j][1]
            if end_merge > start_merge:
                worksheet.merge_cells(start_row=start_merge, start_column=1, end_row=end_merge, end_column=1)
            section_cell = worksheet.cell(row=start_merge, column=1, value=self.state.blocks[i].block_type)
            section_cell.font = Font(bold=True, color="409EFF")
            section_cell.alignment = Alignment(horizontal="center", vertical="center")
            i = j + 1

        worksheet.column_dimensions["C"].width = 40
        worksheet.column_dimensions["D"].width = 40
        return workbook

    @staticmethod
    def _sanitize_output_name(song_name: str) -> str:
        output_name = f"{song_name}_编排表.xlsx"
        for char in ['\\', "/", ":", "*", "?", '"', "<", ">", "|"]:
            output_name = output_name.replace(char, "_")
        return output_name

    def _resolve_output_path(self, output_path: "Path | None" = None) -> Path:
        if output_path is not None:
            return output_path
        if self.default_output_path is not None:
            return self.default_output_path
        return Path(self._sanitize_output_name(self.song_name))

    def _save_workbook_with_retry(self, output_path: "Path | None" = None) -> None:
        target_path = self._resolve_output_path(output_path)
        while True:
            try:
                base = target_path if target_path.exists() else None
                workbook = self._build_workbook(base_path=base)
                target_path.parent.mkdir(parents=True, exist_ok=True)
                workbook.save(target_path)
                self.default_output_path = target_path
                print(f"\n ✨ 成功！文件已保存至: {target_path}")
                return
            except PermissionError:
                input(f"\n ⚠️ 文件 [{target_path}] 被占用，请关闭它后按回车重试...")
